"""
Add the comparison visualization PNG to the Excel file
"""

from openpyxl import load_workbook
from openpyxl.drawing.image import Image

# Load the existing workbook
wb = load_workbook('/home/user/kitchen-demand-lab/Power_Systems_Comparison.xlsx')

# Create new sheet for the visualization
ws = wb.create_sheet('Visualization', 0)  # Insert at the beginning

# Add title
ws['A1'] = 'Revenue & EBITDA Comparison Visualization'
ws['A1'].font = ws['A1'].font.copy(bold=True, size=14)

# Add the PNG image
img = Image('/home/user/kitchen-demand-lab/revenue_ebitda_comparison.png')
# Resize to fit well in Excel (scale to 80%)
img.width = img.width * 0.8
img.height = img.height * 0.8
ws.add_image(img, 'A3')

# Save the workbook
wb.save('/home/user/kitchen-demand-lab/Power_Systems_Comparison.xlsx')

print("Visualization added to Excel file successfully!")
